"""Best-effort XLSX snapshots. Database remains the source of truth."""
import logging
from io import BytesIO
from threading import Lock
import httpx
from openpyxl import Workbook, load_workbook
from sqlalchemy import select, text
from .config import settings
from .database import SessionLocal
from .models import Order, BonusEvent

API = 'https://cloud-api.yandex.net/v1/disk/resources'
lock = Lock()
logger = logging.getLogger(__name__)


def request(client, method, suffix='', **kwargs):
    response = client.request(method, API + suffix, headers={'Authorization': f'OAuth {settings.yandex_disk_token}'}, **kwargs)
    response.raise_for_status()
    return response


def upload_snapshot(client, filename, headers, rows):
    path = '/brand-cabinet/' + filename
    metadata = client.get(API, params={'path': path}, headers={'Authorization': f'OAuth {settings.yandex_disk_token}'})
    if metadata.status_code == 404:
        workbook = Workbook()
    else:
        metadata.raise_for_status()
        href = request(client, 'GET', '/download', params={'path': path}).json()['href']
        response = client.get(href)
        response.raise_for_status()
        workbook = load_workbook(BytesIO(response.content))
    sheet = workbook.active
    sheet.delete_rows(1, sheet.max_row)
    # Reconcile complete DB snapshot, avoiding duplicate rows and stale order statuses.
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = 's'  # Prevent formula injection from customer input.
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    output = BytesIO()
    workbook.save(output)
    href = request(client, 'GET', '/upload', params={'path': path, 'overwrite': 'true'}).json()['href']
    response = client.put(href, content=output.getvalue())
    response.raise_for_status()


def sync_all():
    if not settings.yandex_disk_token:
        return {'status': 'skipped', 'detail': 'YANDEX_DISK_TOKEN не задан'}
    try:
        with lock, SessionLocal() as db, httpx.Client(timeout=30, follow_redirects=True) as client:
            if db.bind.dialect.name == 'postgresql':
                # Serialize exports across backend processes; released with transaction.
                db.execute(text('SELECT pg_advisory_xact_lock(817264)'))
            folder = client.put(API, params={'path': '/brand-cabinet'}, headers={'Authorization': f'OAuth {settings.yandex_disk_token}'})
            if folder.status_code != 409:
                folder.raise_for_status()
            orders = db.scalars(select(Order).order_by(Order.id)).all()
            bonuses = db.scalars(select(BonusEvent).order_by(BonusEvent.id)).all()
            upload_snapshot(client, 'orders.xlsx', ['order_id', 'client_id', 'client_email', 'product_id', 'product_title', 'quantity', 'unit_price', 'total', 'status', 'comment', 'created_at'], [[o.id, o.client_id, o.client.email, o.product_id, o.product.title, o.quantity, str(o.unit_price), str(o.unit_price * o.quantity), o.status, o.comment, o.created_at.isoformat()] for o in orders])
            upload_snapshot(client, 'bonuses.xlsx', ['bonus_event_id', 'client_id', 'client_email', 'amount', 'reason', 'created_at'], [[b.id, b.client_id, b.client.email, str(b.amount), b.reason, b.created_at.isoformat()] for b in bonuses])
        return {'status': 'synced'}
    except Exception as exc:
        # Do not log OAuth tokens or signed download/upload URLs.
        logger.error('Yandex sync failed (%s)', type(exc).__name__)
        return {'status': 'failed', 'detail': 'Данные сохранены в БД; повторите выгрузку позже'}
