import os
import secrets
import tempfile
from datetime import datetime, timedelta, timezone
from concurrent.futures import ThreadPoolExecutor

os.environ['DATABASE_URL'] = 'sqlite:///' + tempfile.mktemp(suffix='.db').replace('\\', '/')
os.environ['SECRET_KEY'] = secrets.token_urlsafe(40)
os.environ['ADMIN_KEY'] = secrets.token_urlsafe(30)
os.environ['YANDEX_DISK_TOKEN'] = ''

import pytest
from fastapi.testclient import TestClient
from jose import jwt
from sqlalchemy import update
from app.main import app
from app.database import Base, engine, SessionLocal
from app.models import Product
from app.config import settings
from app import yandex_sync


@pytest.fixture
def api():
    Base.metadata.drop_all(engine)
    with TestClient(app) as client:
        yield client


def register(api, email='client@example.com'):
    response = api.post('/auth/register', json={'name': 'Анна', 'email': email, 'password': 'safe-password-123'})
    assert response.status_code == 201, response.text
    return {'Authorization': 'Bearer ' + response.json()['access_token']}


def admin():
    return {'admin_key': settings.admin_key}


def product(api, **overrides):
    data = {'title': 'Демо', 'price': '123.45', 'collection': 'Базовая'} | overrides
    response = api.post('/admin/products', params=admin(), json=data)
    assert response.status_code == 201, response.text
    return response.json()['id']


def test_full_journey(api):
    headers = register(api)
    assert api.get('/health').status_code == 200
    assert api.get('/me', headers=headers).json()['bonus_balance'] == '100.00'
    assert api.post('/auth/login', json={'email': 'CLIENT@example.com', 'password': 'safe-password-123'}).status_code == 200
    assert api.patch('/me', headers=headers, json={'name': 'Анна Тест', 'phone': '+7 900 000 00 00'}).json()['name'] == 'Анна Тест'
    pid = product(api)
    assert len(api.get('/products?collection=Базовая').json()) == 1
    assert api.get('/products?collection=Нет').json() == []
    response = api.post('/orders', headers=headers, json={'product_id': pid, 'quantity': 2})
    assert response.status_code == 201, response.text
    oid = response.json()['id']
    with SessionLocal() as db:
        db.execute(update(Product).where(Product.id == pid).values(price=999))
        db.commit()
    assert api.get('/me/orders', headers=headers).json()[0]['unit_price'] == '123.45'
    assert api.patch(f'/admin/orders/{oid}/status', params=admin(), json={'status': 'done'}).status_code == 409
    for status in ['processing', 'confirmed', 'shipped', 'done', 'done']:
        assert api.patch(f'/admin/orders/{oid}/status', params=admin(), json={'status': status}).status_code == 200
    assert api.get('/me', headers=headers).json()['bonus_balance'] == '112.35'
    assert len(api.get('/me/bonuses', headers=headers).json()) == 2
    assert api.patch(f'/admin/orders/{oid}/status', params=admin(), json={'status': 'cancelled'}).status_code == 409
    assert api.post('/admin/sync-all', params=admin()).json()['status'] == 'skipped'


def test_auth_validation_and_isolation(api):
    headers = register(api)
    assert api.post('/auth/register', json={'name': 'X', 'email': 'CLIENT@example.com', 'password': 'safe-password-123'}).status_code == 409
    assert api.post('/auth/login', json={'email': 'client@example.com', 'password': 'wrong-password'}).status_code == 401
    assert api.get('/me').status_code == 401
    assert api.get('/me', headers={'Authorization': 'Bearer invalid'}).status_code == 401
    expired = jwt.encode({'sub': '1', 'exp': datetime.now(timezone.utc) - timedelta(hours=1)}, settings.secret_key, algorithm='HS256')
    assert api.get('/me', headers={'Authorization': 'Bearer ' + expired}).status_code == 401
    assert api.get('/admin/orders').status_code == 403
    assert api.post('/orders', headers=headers, json={'product_id': 999, 'quantity': 1}).status_code == 404
    pid = product(api, in_stock=False)
    assert api.post('/orders', headers=headers, json={'product_id': pid, 'quantity': 1}).status_code == 409
    assert api.post('/orders', headers=headers, json={'product_id': pid, 'quantity': 0}).status_code == 422
    assert api.post('/orders', headers=headers, json={'product_id': pid, 'quantity': 1.5}).status_code == 422
    pid = product(api)
    api.post('/orders', headers=headers, json={'product_id': pid, 'quantity': 1})
    other = register(api, 'other@example.com')
    assert api.get('/me/orders', headers=other).json() == []
    assert api.patch('/me', headers=other, json={'name': 'X', 'bonus_balance': 999}).status_code == 422


def test_parallel_cashback(api):
    headers = register(api)
    pid = product(api)
    oid = api.post('/orders', headers=headers, json={'product_id': pid, 'quantity': 1}).json()['id']
    for status in ['processing', 'confirmed', 'shipped']:
        api.patch(f'/admin/orders/{oid}/status', params=admin(), json={'status': status})
    def complete(_):
        return api.patch(f'/admin/orders/{oid}/status', params=admin(), json={'status': 'done'}).status_code
    with ThreadPoolExecutor(max_workers=4) as executor:
        assert all(code in (200, 409) for code in executor.map(complete, range(4)))
    assert api.get('/me', headers=headers).json()['bonus_balance'] == '106.17'
    assert len(api.get('/me/bonuses', headers=headers).json()) == 2


def test_praktium_pending_price_requires_quote(api):
    headers = register(api)
    pid = product(api, title='Массажёр для шеи и плеч', price='0.00', collection='Здоровье')
    order = api.post('/orders', headers=headers, json={'product_id': pid, 'quantity': 2})
    assert order.status_code == 201
    oid = order.json()['id']
    assert order.json()['unit_price'] == '0.00'
    for status in ['processing', 'confirmed', 'shipped']:
        assert api.patch(f'/admin/orders/{oid}/status', params=admin(), json={'status': status}).status_code == 200
    assert api.patch(f'/admin/orders/{oid}/status', params=admin(), json={'status': 'done'}).status_code == 409
    assert api.patch(f'/admin/orders/{oid}/price', params=admin(), json={'unit_price': '0'}).status_code == 422
    quote = api.patch(f'/admin/orders/{oid}/price', params=admin(), json={'unit_price': '2500.00'})
    assert quote.status_code == 200
    assert quote.json()['unit_price'] == '2500.00'
    assert api.patch(f'/admin/orders/{oid}/status', params=admin(), json={'status': 'done'}).status_code == 200
    assert api.get('/me', headers=headers).json()['bonus_balance'] == '350.00'
    assert api.patch(f'/admin/orders/{oid}/price', params=admin(), json={'unit_price': '2600.00'}).status_code == 409


def test_sync_failure_does_not_break_registration(api, monkeypatch):
    monkeypatch.setattr(settings, 'yandex_disk_token', 'test-only')
    def fail(*args, **kwargs):
        raise RuntimeError('Simulated outage')
    monkeypatch.setattr(yandex_sync.httpx, 'Client', fail)
    headers = register(api)
    assert api.get('/me', headers=headers).status_code == 200
    assert api.post('/admin/sync-all', params=admin()).json()['status'] == 'failed'


def test_partial_profile_and_password_bytes(api):
    headers = register(api)
    assert api.patch('/me', headers=headers, json={'phone': '123'}).json()['name'] == 'Анна'
    assert api.patch('/me', headers=headers, json={'name': None}).status_code == 422
    assert api.post('/auth/register', json={'name': 'X', 'email': 'long@example.com', 'password': 'я' * 40}).status_code == 422


def test_yandex_snapshot_upsert_and_formula_safety(api, monkeypatch):
    from io import BytesIO
    from openpyxl import load_workbook
    import httpx
    headers = register(api)
    pid = product(api)
    api.post('/orders', headers=headers, json={'product_id': pid, 'quantity': 2, 'comment': '=1+1'})
    uploaded = {}
    def handle(request):
        path = request.url.path
        if request.url.host == 'storage.test':
            filename = path.rsplit('/', 1)[-1]
            assert 'authorization' not in request.headers
            if request.method == 'PUT':
                uploaded[filename] = request.content
                return httpx.Response(201)
            return httpx.Response(200, content=uploaded[filename])
        assert request.headers['authorization'] == 'OAuth test-only'
        if request.method == 'PUT':
            return httpx.Response(409)
        filename = request.url.params.get('path', '').rsplit('/', 1)[-1]
        if path.endswith('/upload'):
            return httpx.Response(200, json={'href': 'https://storage.test/' + filename})
        if path.endswith('/download'):
            return httpx.Response(200, json={'href': 'https://storage.test/' + filename})
        return httpx.Response(200 if filename in uploaded else 404, json={})
    original_client = httpx.Client
    monkeypatch.setattr(settings, 'yandex_disk_token', 'test-only')
    monkeypatch.setattr(yandex_sync.httpx, 'Client', lambda **kwargs: original_client(transport=httpx.MockTransport(handle), **kwargs))
    for _ in range(2):
        assert yandex_sync.sync_all()['status'] == 'synced'
    orders = load_workbook(BytesIO(uploaded['orders.xlsx'])).active
    bonuses = load_workbook(BytesIO(uploaded['bonuses.xlsx'])).active
    assert orders.max_row == 2
    assert bonuses.max_row == 2
    assert orders['J2'].value == '=1+1'
    assert orders['J2'].data_type == 's'
    assert orders['G2'].value == '123.45'
